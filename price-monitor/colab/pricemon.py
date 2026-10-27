"""
pricemon — монитор цен конкурентов (Python-версия расширения «Монитор цен конкурентов»).

Работает в Google Colab, Jupyter или обычном Python 3.9+.
  • проверяет ссылки по одной, со случайной паузой (по умолчанию 5–10 с) и чередует сайты;
  • один и тот же сайт — не чаще чем раз в N секунд;
  • читает цену: CSS-селектор → якорь (артикул/название на странице-списке) → JSON-LD → microdata → meta
    → цена рядом с артикулом → эвристика по классам «price» (без старой цены, рассрочки, «похожих товаров»);
  • капча / «слишком много запросов» / 403 / 429 — сайт пропускается до конца прогона, защиту НЕ обходим;
  • 5xx и сетевые ошибки — один повтор в конце очереди;
  • отчёт: Цена товара | Разница в цене в % | Магазин демпинг | Ссылка на магазин конкурента.

Формат списка позиций совпадает с «Выгрузить позиции» в расширении:
  Наименование ; Артикул ; Наша цена ; Ссылка [; Якорь ; CSS-селектор]
"""
from __future__ import annotations

import csv
import io
import json
import random
import re
import time
from dataclasses import dataclass, field, asdict
from datetime import datetime
from typing import Callable, Iterable, Optional
from urllib.parse import urlparse

__version__ = "1.0.0"

# --------------------------------------------------------------------------------------------
# Настройки
# --------------------------------------------------------------------------------------------
DEFAULT_EXCLUDED = ["comfysan.ru", "comfysan-shop.ru", "polotenhik.ru", "sancomfgroup.ru", "sancomf.ru"]
# Один фиксированный User-Agent, который честно называет себя монитором цен (без ротации и маскировки).
DEFAULT_USER_AGENT = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/124.0 Safari/537.36 PriceMonitor/" + __version__
)


@dataclass
class Settings:
    min_pause: float = 5.0           # пауза между запросами, секунд (случайная в диапазоне)
    max_pause: float = 10.0
    domain_gap: float = 0.0          # один сайт — не чаще, чем раз в N секунд (0 — только общая пауза)
    timeout: float = 30.0            # таймаут запроса, секунд
    retry_failed: bool = True        # повторить 5xx/сетевые ошибки один раз в конце очереди
    dump_threshold_pct: float = 0.0  # демпинг, если цена ниже нашей больше чем на N %
    excluded_domains: list = field(default_factory=lambda: list(DEFAULT_EXCLUDED))
    user_agent: str = DEFAULT_USER_AGENT
    use_browser_for: list = field(default_factory=list)  # домены, которые открывать через Playwright (JS-сайты)

    def normalized(self) -> "Settings":
        s = Settings(**asdict(self))
        s.min_pause = max(0.0, float(s.min_pause))
        s.max_pause = max(s.min_pause, float(s.max_pause))
        s.domain_gap = max(0.0, float(s.domain_gap))
        if isinstance(s.excluded_domains, str):
            s.excluded_domains = re.split(r"[\s,;]+", s.excluded_domains)
        s.excluded_domains = [host_of(d) or d.strip().lower() for d in s.excluded_domains if str(d).strip()]
        if isinstance(s.use_browser_for, str):
            s.use_browser_for = re.split(r"[\s,;]+", s.use_browser_for)
        s.use_browser_for = [site_of(d) or d.strip().lower() for d in s.use_browser_for if str(d).strip()]
        return s


# --------------------------------------------------------------------------------------------
# Домены, цены
# --------------------------------------------------------------------------------------------
def host_of(url: str) -> str:
    s = str(url or "").strip()
    if not s:
        return ""
    if not re.match(r"^[a-z][a-z0-9+.-]*://", s, re.I):
        s = "https://" + s
    try:
        h = (urlparse(s).hostname or "").lower()
    except ValueError:
        return ""
    return h[4:] if h.startswith("www.") else h


def site_of(url: str) -> str:
    """msk.santehnica.ru и santehnica.ru — один сайт."""
    h = host_of(url)
    if not h or re.match(r"^\d+\.\d+\.\d+\.\d+$", h):
        return h
    p = h.split(".")
    if len(p) <= 2:
        return h
    n = 3 if (p[-2] in ("com", "net", "org", "co", "gov", "edu", "msk", "spb") and len(p[-1]) == 2) else 2
    return ".".join(p[-n:])


def is_excluded(url: str, excluded: Iterable[str]) -> bool:
    h = host_of(url)
    return bool(h) and any(h == d or h.endswith("." + d) for d in excluded)


_PRICE_RE = re.compile(r"(\d{1,3}(?:[ \u00a0\u2009\u202f'](?=\d{3})\d{3})+|\d+)(?:[.,](\d{1,2}))?(?!\d)")


def parse_price(text, lo: float = 50, hi: float = 50_000_000) -> Optional[float]:
    """Первая правдоподобная цена: «23 715 ₽», «23&nbsp;715 руб.», «23715.00», «от 1 999,90». «-15%» пропускается."""
    if isinstance(text, (int, float)):
        return float(text) if text > 0 else None
    s = "" if text is None else str(text)
    for m in _PRICE_RE.finditer(s):
        if re.match(r"^\s*%", s[m.end():m.end() + 3]):
            continue
        v = float(re.sub(r"\D", "", m.group(1)) + ("." + m.group(2) if m.group(2) else ""))
        if lo <= v <= hi:
            return round(v, 2)
    return None


def deviation(competitor, ours) -> Optional[float]:
    """((конкурент − наша) / наша) × 100, 2 знака."""
    try:
        c, o = float(competitor), float(ours)
    except (TypeError, ValueError):
        return None
    if not (o > 0 and c > 0):
        return None
    return round((c - o) / o * 100, 2)


def is_dumping(competitor, ours, threshold_pct: float = 0.0) -> Optional[bool]:
    d = deviation(competitor, ours)
    return None if d is None else d < -float(threshold_pct or 0) - 1e-9


def fmt_money(v) -> str:
    if v is None:
        return "—"
    s = f"{v:,.2f}" if v % 1 else f"{v:,.0f}"
    return s.replace(",", "\u00a0").replace(".", ",") + "\u00a0₽"


def fmt_pct(d) -> str:
    if d is None:
        return "—"
    sign = "+" if d > 0 else ("−" if d < 0 else "")
    return f"{sign}{abs(d):.2f}%".replace(".", ",")


# --------------------------------------------------------------------------------------------
# Позиции
# --------------------------------------------------------------------------------------------
@dataclass
class Link:
    url: str
    anchor: str = ""
    selector: str = ""
    enabled: bool = True


@dataclass
class Position:
    name: str
    article: str = ""
    our_price: float = 0.0
    links: list = field(default_factory=list)


SAMPLE_POSITIONS = [
    Position("Comfysan Альто-К (Alto-K) EC-3 120/12, оружейная сталь, гальваника", "018160", 27900, [
        Link("https://www.santehtop.ru/comfysan-alto-k-018160/"),
        Link("https://santehpremium.ru/product/polotencesushitel-elektricheskiy-s-t-comfysan-alto-k-ec-3-120-15-oruzheynaya-stal-nr-018160/"),
        Link("https://www.santehnica.ru/product/1644811.html"),
        Link("https://santeh-crystal.ru/polotentsesushitel-elektricheskij-st-comfysan-alto-k-ec-2-12010-belyj-pravyj-ten",
             anchor="EC-3 120/12, оружейная сталь, гальвани"),
    ]),
    Position("Comfysan Альто-К (Alto-K) EC-4 120/18, оружейная сталь, гальваника", "018184", 34100, [
        Link("https://www.santehtop.ru/comfysan-alto-k-018184/"),
        Link("https://santehpremium.ru/product/polotencesushitel-elektricheskiy-s-t-comfysan-alto-k-ec-4-120-15-oruzheynaya-stal-nr-018184/"),
        Link("https://santeh-crystal.ru/polotentsesushitel-elektricheskij-st-comfysan-alto-k-ec-2-12010-belyj-pravyj-ten",
             anchor="EC-4 120/18, оружейная сталь, гальваника"),
        Link("https://www.vseinstrumenti.ru/tag-page/polotentsesushiteli-elektricheskie-vertikalnye-674912/", anchor="018184"),
    ]),
    Position("Comfysan Альто-К (Alto-K) EC-4 120/18, золото, гальваника", "018443", 34100, [
        Link("https://www.santehtop.ru/comfysan-alto-k-018443/"),
        Link("https://santehpremium.ru/product/polotencesushitel-elektricheskiy-s-t-comfysan-alto-k-ec-4-120-15-zoloto-galvanika-018443/"),
        Link("https://santeh-crystal.ru/polotentsesushitel-elektricheskij-st-comfysan-alto-k-ec-2-12010-belyj-pravyj-ten",
             anchor="EC-4 120/18, золото, гальваника"),
        Link("https://www.vseinstrumenti.ru/tag-page/polotentsesushiteli-elektricheskie-vertikalnye-674912/", anchor="018443"),
    ]),
    Position("Comfysan Стелла (Stella) EC-6 70/40, оружейная сталь, гальваника", "018191", 22500, [
        Link("https://www.santehtop.ru/comfysan-stella-018191/"),
        Link("https://www.santehnica.ru/product/1650561.html"),
    ]),
]


def _sniff_rows(text: str) -> list:
    text = text.lstrip("\ufeff")
    first = text.splitlines()[0] if text else ""
    delim = "\t" if "\t" in first else (";" if first.count(";") >= first.count(",") else ",")
    return [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim) if any(c.strip() for c in r)]


def positions_from_rows(rows: list) -> list:
    """Строки таблицы → позиции. Шапка необязательна; строки с одинаковым артикулом объединяются."""
    rows = [[("" if c is None else str(c)).strip() for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    idx = {"name": 0, "article": 1, "price": 2, "url": 3, "anchor": 4, "selector": 5}
    head = [c.lower() for c in rows[0]]
    looks_head = any(re.search(r"наимен|товар|назв|артик|ссылк|url|цен", c) for c in head) and \
        not any(c.startswith("http") or re.fullmatch(r"[\d\s.,₽р]+", c) for c in head)
    if looks_head:
        def find(rx):
            for i, c in enumerate(head):
                if re.search(rx, c):
                    return i
            return -1
        idx = {"name": find(r"наимен|товар|назв|name"), "article": find(r"артик|код|sku"),
               "price": find(r"наша|ррц|цена|price"), "url": find(r"ссылк|url|адрес|link"),
               "anchor": find(r"якор|anchor|текст"), "selector": find(r"селект|selector|css")}
        rows = rows[1:]
    out, by_key = [], {}
    for r in rows:
        g = lambda k: r[idx[k]] if 0 <= idx[k] < len(r) else ""  # noqa: E731
        name, article, url = g("name"), g("article"), g("url")
        if not url:
            url = next((c for c in r if c.lower().startswith("http")), "")
        if article.endswith(".0") and article[:-2].isdigit():  # артикул из Excel как число
            article = article[:-2]
        if not name and not article:
            continue
        key = (article or name).lower()
        p = by_key.get(key)
        if p is None:
            p = Position(name=name, article=article, our_price=parse_price(g("price")) or 0.0)
            by_key[key] = p
            out.append(p)
        else:
            p.name = p.name or name
            p.our_price = parse_price(g("price")) or p.our_price
        enabled = True
        if url.startswith("#"):
            enabled, url = False, url.lstrip("# ").strip()
        if url.lower().startswith("http") and not any(l.url == url and l.anchor == g("anchor") for l in p.links):
            p.links.append(Link(url=url, anchor=g("anchor"), selector=g("selector"), enabled=enabled))
    return out


def load_positions(source) -> list:
    """CSV/TSV-текст, путь к .csv/.xlsx/.xls или pandas.DataFrame."""
    try:
        import pandas as pd  # noqa: F401
    except ImportError:  # pragma: no cover
        pd = None
    if pd is not None and hasattr(source, "values") and hasattr(source, "columns"):
        return positions_from_rows([list(map(str, source.columns))] + source.astype(str).values.tolist())
    s = str(source)
    if "\n" not in s and re.search(r"\.(xlsx|xlsm|xls)$", s, re.I):
        df = pd.read_excel(s, header=None, dtype=str).fillna("")
        return positions_from_rows(df.values.tolist())
    if "\n" not in s and re.search(r"\.(csv|tsv|txt)$", s, re.I):
        raw = open(s, "rb").read()
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError:
            text = raw.decode("cp1251")  # CSV из русского Excel
        return positions_from_rows(_sniff_rows(text))
    return positions_from_rows(_sniff_rows(s))


def positions_to_csv(positions: list) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["Наименование", "Артикул", "Наша цена", "Ссылка", "Якорь", "CSS-селектор"])
    for p in positions:
        for l in (p.links or [Link("")]):
            w.writerow([p.name, p.article, int(p.our_price) if p.our_price else "",
                        ("" if l.enabled else "# ") + l.url, l.anchor, l.selector])
    return "\ufeff" + buf.getvalue()


# --------------------------------------------------------------------------------------------
# Извлечение цены из HTML
# --------------------------------------------------------------------------------------------
BLOCK_RE = re.compile(
    r"captcha|капч|cf-chl|just a moment|checking your browser|проверка браузера|проверяем браузер|ddos-guard|"
    r"превышени[ея] (?:числа|количества|лимита) запросов|слишком много запросов|too many requests|access denied|"
    r"доступ (?:запрещ|огранич)|подтвердите,? что вы не робот|вы не робот|are you a robot|attention required|"
    r"security check|smartcaptcha", re.I)
NOTFOUND_RE = re.compile(r"(^|\D)404(\D|$)|страница не найдена|page not found|товар не найден|не существует|снят с продажи", re.I)
BAD_CLASS_RE = re.compile(r"old|prev|was|cross|strike|through|economy|save|saving|discount|diff|credit|month|bonus|install|"
                          r"rassroch|compare|cashback|delivery|dostav|cart-total|basket|minicart|header|footer|filter|range|"
                          r"slider-price|price-?filter", re.I)
RELATED_RE = re.compile(r"similar|related|recommend|viewed|analog|accessor|carousel|swiper|slick|upsell|cross-?sell|also|"
                        r"похож|together|bundle|mini-?card", re.I)
PRICEISH_RE = re.compile(r"price|cost|cena|tsena|sum|amount", re.I)
BAD_TEXT_RE = re.compile(r"мес|month|бонус|балл|кешбэк|cashback|скидк|эконом|выгод", re.I)


def norm(s) -> str:
    return re.sub(r"[\u00a0\u2009\u202f\s]+", " ", str(s or "").lower().replace("ё", "е")).strip()


def _cls(el) -> str:
    if not getattr(el, "attrs", None):
        return ""
    c = el.get("class") or []
    c = " ".join(c) if isinstance(c, list) else str(c)
    return " ".join([c, el.get("id") or "", el.get("itemprop") or "", el.get("data-qa") or el.get("data-testid") or ""])


def _ancestors(el, depth):
    e, i = el, 0
    while e is not None and getattr(e, "name", None) and i < depth:
        yield e
        e, i = e.parent, i + 1


def _hidden(el) -> bool:
    for e in _ancestors(el, 12):
        if e.name in ("[document]",):
            break
        if e.has_attr("hidden") or e.get("aria-hidden") == "true":
            return True
        st = (e.get("style") or "").replace(" ", "").lower()
        if "display:none" in st or "visibility:hidden" in st:
            return True
    return False


def _struck(el) -> bool:
    for e in _ancestors(el, 4):
        if e.name in ("s", "del", "strike"):
            return True
        if "line-through" in (e.get("style") or ""):
            return True
    return False


def _bad_ancestor(el, rx, depth) -> bool:
    for e in _ancestors(el, depth):
        if e.name in ("body", "main", "[document]"):
            return False
        if e.name in ("header", "footer", "nav", "aside"):
            return True
        if rx.search(_cls(e)):
            return True
    return False


def _text_of(el) -> str:
    return el.get("content") or el.get_text(" ", strip=True)


def _is_price_el(tag) -> bool:
    if not getattr(tag, "attrs", None):
        return False
    if tag.get("itemprop") == "price" or tag.has_attr("data-price"):
        return True
    return bool(PRICEISH_RE.search(" ".join(tag.get("class") or []) + " " + (tag.get("id") or "")))


def price_candidates(scope, allow_related=False) -> list:
    out = []
    for el in scope.find_all(_is_price_el):
        c = _cls(el)
        inner = next((x for x in el.find_all(_is_price_el) if x is not el), None)
        if inner is not None and parse_price(_text_of(inner)) is not None and el.get("itemprop") != "price":
            continue
        if BAD_CLASS_RE.search(c):
            continue
        txt = _text_of(el) if el.get("itemprop") == "price" else (el.get("data-price") or el.get_text(" ", strip=True))
        if len(txt) > 80 or BAD_TEXT_RE.search(txt):
            continue
        v = parse_price(txt)
        if v is None or _struck(el) or _hidden(el):
            continue
        if _bad_ancestor(el, BAD_CLASS_RE, 3):
            continue
        if not allow_related and _bad_ancestor(el, RELATED_RE, 14):
            continue
        out.append((el, v))
    return out


def _common_depth(a, b, maxd) -> float:
    if a is None or b is None:
        return float("inf")
    bpar = set(id(x) for x in _ancestors(b, 60))
    for i, e in enumerate(_ancestors(a, maxd)):
        if id(e) in bpar:
            return i
    return float("inf")


def best_candidate(soup, cands):
    if not cands:
        return None
    h1 = soup.find("h1")
    best, best_score = None, float("-inf")
    for i, (el, v) in enumerate(cands):
        near = _common_depth(h1, el, 8)
        score = (8 - near) * 15 if near < 8 else 0
        if re.search(r"current|actual|new|final|sale|main", _cls(el), re.I):
            score += 30
        if el.name in ("b", "strong") or re.search(r"big|large|lg|xl", _cls(el), re.I):
            score += 10
        score -= i * 0.01
        if score > best_score:
            best, best_score = (el, v), score
    return best


def _jsonld_products(soup) -> list:
    out = []

    def visit(n):
        if isinstance(n, list):
            for x in n:
                visit(x)
            return
        if not isinstance(n, dict):
            return
        t = n.get("@type")
        if t in ("Product", "ProductGroup") or (isinstance(t, list) and "Product" in t):
            out.append(n)
        for k in ("@graph", "itemListElement", "hasVariant"):
            if k in n:
                visit(n[k])
        if isinstance(n.get("item"), dict):
            visit(n["item"])

    for sc in soup.find_all("script", attrs={"type": "application/ld+json"}):
        try:
            visit(json.loads(re.sub(r"[\x00-\x1f]+", " ", sc.string or sc.get_text() or "")))
        except (ValueError, TypeError):
            pass
    return out


def _offer_price(p):
    offers = p.get("offers")
    if not offers:
        return None
    lst = offers if isinstance(offers, list) else [offers]
    prices, avail = [], None
    for o in lst:
        if not isinstance(o, dict):
            continue
        raw = o.get("price", o.get("lowPrice", (o.get("priceSpecification") or {}).get("price") if isinstance(o.get("priceSpecification"), dict) else None))
        v = parse_price(raw)
        if v is not None:
            prices.append(v)
        avail = avail or o.get("availability")
    return (min(prices), avail) if prices else None


def _availability(s) -> Optional[str]:
    s = norm(s)
    if not s:
        return None
    if re.search(r"outofstock|нет в наличии|нет на складе|отсутствует|распродан|снят с", s):
        return "Нет в наличии"
    if re.search(r"preorder|backorder|под заказ|предзаказ|ожидается", s):
        return "Под заказ"
    if re.search(r"instock|limitedavailability|в наличии|есть на складе|в магазине|на складе", s):
        return "В наличии"
    return None


def _nearby_availability(el) -> Optional[str]:
    for e in _ancestors(el, 6):
        a = _availability(e.get_text(" ", strip=True)[:3000])
        if a:
            return a
    return None


def _anchor_elements(soup, anchor_n):
    body = soup.body or soup
    out = []
    for t in body.find_all(string=True):
        if t.parent is not None and t.parent.name in ("script", "style", "noscript", "template"):
            continue
        if anchor_n in norm(t):
            out.append(t.parent)
    if not out:
        for el in body.find_all(["a", "h1", "h2", "h3", "h4"]):
            txt = el.get_text(" ", strip=True)
            if len(txt) < 400 and anchor_n in norm(txt):
                out.append(el)
    if re.fullmatch(r"[\w-]{4,}", anchor_n):
        out += [a for a in body.find_all("a", href=True) if anchor_n in norm(a["href"])]
    return out


def _price_near_anchor(soup, anchor_n, allow_related):
    for a in _anchor_elements(soup, anchor_n):
        if _hidden(a):
            continue
        for depth, e in enumerate(_ancestors(a, 9)):
            if e.name in ("body", "[document]"):
                break
            c = price_candidates(e, allow_related=allow_related)
            if c:
                b = c[0] if len(c) == 1 else best_candidate(soup, c)
                return b[1], e
    return None


def extract(html: str, url: str = "", article: str = "", anchor: str = "", selector: str = "") -> dict:
    """Цена и признаки со страницы. Возвращает dict: price, method, availability, blocked, not_found, article_found, title."""
    from bs4 import BeautifulSoup
    try:
        soup = BeautifulSoup(html or "", "lxml")
    except Exception:  # lxml не установлен
        soup = BeautifulSoup(html or "", "html.parser")
    body = soup.body or soup
    for junk in body.find_all(["script", "style", "noscript", "template"]):
        if junk.get("type") != "application/ld+json":
            junk.extract()
    body_text = body.get_text(" ", strip=True)
    title = (soup.title.get_text(strip=True) if soup.title else "")
    h1 = soup.find("h1")
    art, anc = norm(article), norm(anchor)
    res = {"url": url, "title": h1.get_text(" ", strip=True) if h1 else title, "price": None, "method": None,
           "availability": None, "blocked": None, "not_found": False,
           "article_found": (art in norm(body_text) or art in norm(url)) if art else None,
           "anchor_found": None, "text_length": len(body_text)}

    def done(price, method, el=None, avail=None):
        res.update(price=price, method=method, availability=_availability(avail) or (_nearby_availability(el) if el is not None else None))
        return finish()

    def finish():
        if res["price"] is None:
            probe = title + " \n" + body_text[:4000]
            challenge = soup.select_one("#challenge-form, #cf-challenge-running, .cf-browser-verification, "
                                        "iframe[src*=captcha], iframe[src*=challenge], #captcha, .captcha, [class*=smart-captcha]")
            if BLOCK_RE.search(title) or ((challenge is not None or BLOCK_RE.search(probe)) and len(body_text) < 6000):
                m = BLOCK_RE.search(probe)
                res["blocked"] = m.group(0) if m else "капча"
            elif NOTFOUND_RE.search(title) or (NOTFOUND_RE.search(body_text[:1500]) and len(body_text) < 8000):
                res["not_found"] = True
        return res

    if selector:
        try:
            el = soup.select_one(selector)
            v = parse_price(_text_of(el)) if el is not None else None
            if v is not None:
                return done(v, "selector", el)
        except Exception as e:  # кривой селектор
            res["selector_error"] = str(e)
    if anc:
        near = _price_near_anchor(soup, anc, allow_related=True)
        if near:
            return done(near[0], "anchor", near[1])
        res["anchor_found"] = bool(_anchor_elements(soup, anc))
        return finish()
    prods = [(p, op) for p in _jsonld_products(soup) for op in [_offer_price(p)] if op]
    if prods:
        pick = prods[0] if len(prods) == 1 else next(
            (x for x in prods if art and art in norm(" ".join(str(x[0].get(k, "")) for k in ("sku", "mpn", "productID", "name", "url")))), None)
        if pick is None and len({x[1][0] for x in prods}) == 1:
            pick = prods[0]
        if pick:
            return done(pick[1][0], "json-ld", None, pick[1][1])
    md = []
    for pe in soup.select('[itemtype*="schema.org/Product"]'):
        if _bad_ancestor(pe.parent, RELATED_RE, 14) if pe.parent is not None else False:
            continue
        pr = pe.select_one('[itemprop="price"], [itemprop="lowPrice"]')
        v = parse_price(_text_of(pr)) if pr is not None else None
        if v is not None:
            av = pe.select_one('[itemprop="availability"]')
            md.append((pe, v, (av.get("href") or av.get("content") or av.get_text()) if av is not None else None))
    if md:
        mp = md[0] if len(md) == 1 else next((x for x in md if art and art in norm(x[0].get_text(" "))), None)
        if mp:
            return done(mp[1], "microdata", mp[0], mp[2])
    meta = soup.select_one('meta[property="product:price:amount"], meta[property="og:price:amount"], meta[itemprop="price"]')
    if meta is not None and parse_price(meta.get("content")) is not None:
        return done(parse_price(meta.get("content")), "meta")
    if art and res["article_found"]:
        near = _price_near_anchor(soup, art, allow_related=False)
        if near:
            return done(near[0], "article", near[1])
    b = best_candidate(soup, price_candidates(body))
    if b:
        return done(b[1], "heuristic", b[0])
    return finish()


# --------------------------------------------------------------------------------------------
# Загрузка страниц
# --------------------------------------------------------------------------------------------
class Fetcher:
    """requests.Session с одним фиксированным UA; для доменов из use_browser_for — Playwright (если установлен)."""

    def __init__(self, settings: Settings):
        import requests
        self.s = settings
        self.session = requests.Session()
        self.session.headers.update({
            "User-Agent": settings.user_agent,
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "ru-RU,ru;q=0.9,en;q=0.5",
        })
        self._pw = None
        self._browser = None

    def get(self, url: str):
        """→ (status_code, html, final_url)."""
        if site_of(url) in self.s.use_browser_for:
            return self._get_browser(url)
        r = self.session.get(url, timeout=self.s.timeout, allow_redirects=True)
        if not r.encoding or r.encoding.lower() == "iso-8859-1":
            r.encoding = r.apparent_encoding
        return r.status_code, r.text, r.url

    def _get_browser(self, url: str):
        if self._browser is None:
            from playwright.sync_api import sync_playwright  # pip install playwright && playwright install chromium
            self._pw = sync_playwright().start()
            self._browser = self._pw.chromium.launch()
        page = self._browser.new_page(user_agent=self.s.user_agent, locale="ru-RU")
        try:
            resp = page.goto(url, timeout=self.s.timeout * 1000, wait_until="domcontentloaded")
            page.wait_for_timeout(4000)  # дать JS дорисовать цену
            return (resp.status if resp else 0), page.content(), page.url
        finally:
            page.close()

    def close(self):
        try:
            if self._browser:
                self._browser.close()
            if self._pw:
                self._pw.stop()
        except Exception:
            pass


# --------------------------------------------------------------------------------------------
# Прогон
# --------------------------------------------------------------------------------------------
STATUS_TEXT = {"ok": "OK", "noprice": "Цена не найдена", "blocked": "Сайт ограничил доступ", "error": "Ошибка",
               "notfound": "Страница не найдена", "skipped": "Пропущено (сайт ограничил доступ)", "excluded": "Исключён"}


def build_queue(positions: list, settings: Settings, rnd: random.Random = None) -> list:
    """Задачи (позиция, ссылка), перемешаны и чередуются по сайтам; одинаковые URL открываются один раз."""
    rnd = rnd or random.Random()
    tasks = [(p, l) for p in positions for l in p.links if l.enabled and not is_excluded(l.url, settings.excluded_domains)]
    rnd.shuffle(tasks)
    by_url, order_urls = {}, []
    for p, l in tasks:
        if l.url not in by_url:
            by_url[l.url] = []
            order_urls.append(l.url)
        by_url[l.url].append((p, l))
    by_site = {}
    for u in order_urls:
        by_site.setdefault(site_of(u), []).append(u)
    sites = sorted(by_site, key=lambda s: -len(by_site[s]))
    queue = []
    while any(by_site[s] for s in sites):
        for s in sites:
            if by_site[s]:
                u = by_site[s].pop(0)
                queue.append((u, by_url[u]))
    return queue


@dataclass
class Result:
    position: Position
    link: Link
    status: str = "error"
    price: Optional[float] = None
    availability: str = ""
    article_found: Optional[bool] = None
    method: str = ""
    error: str = ""
    final_url: str = ""
    checked_at: str = ""


def run(positions: list, settings: Settings = None, log: Callable[[str], None] = print,
        sleep: Callable[[float], None] = time.sleep, fetcher=None, rnd: random.Random = None,
        stop=None) -> list:
    """Один круг проверки. Возвращает список Result (по одному на каждую ссылку каждой позиции).
    stop — threading.Event: если установлен, оставшиеся страницы не проверяются."""
    st = (settings or Settings()).normalized()
    rnd = rnd or random.Random()
    own_fetcher = fetcher is None
    fetcher = fetcher or Fetcher(st)
    queue = build_queue(positions, st, rnd)
    results, blocked_sites, last_hit, retry = {}, {}, {}, []
    total = len(queue)
    log(f"В очереди {total} страниц, пауза {st.min_pause:g}–{st.max_pause:g} с"
        + (f", один сайт не чаще раза в {st.domain_gap:g} с" if st.domain_gap else "") + ".")
    t_start = time.time()
    first = True

    def process(url, items, attempt):
        nonlocal first
        site = site_of(url)
        if site in blocked_sites:
            for p, l in items:
                results[id(l)] = Result(p, l, status="skipped", error=f"{site} ранее показал защиту: {blocked_sites[site]}",
                                        final_url=url, checked_at=_now())
            log(f"  пропуск {url} — сайт {site} уже ограничил доступ")
            return
        pause = 0.0 if first else rnd.uniform(st.min_pause, st.max_pause)
        if st.domain_gap and site in last_hit:
            pause = max(pause, st.domain_gap - (time.time() - last_hit[site]))
        if pause > 0:
            sleep(pause)
        first = False
        stamp = _now()
        try:
            code, html, final = fetcher.get(url)
            err = None
        except Exception as e:  # сеть, таймаут, DNS
            code, html, final, err = 0, "", url, f"{type(e).__name__}: {e}"
        last_hit[site] = time.time()
        if err or code >= 500:
            msg = err or f"HTTP {code}"
            if attempt == 0 and st.retry_failed:
                retry.append((url, items))
                log(f"{stamp} [{pause:4.1f} с] {host_of(url)}: {msg} — повторю в конце очереди")
                return
            for p, l in items:
                results[id(l)] = Result(p, l, status="error", error=msg, final_url=final, checked_at=stamp)
            log(f"{stamp} [{pause:4.1f} с] {host_of(url)}: {msg}")
            return
        for p, l in items:
            r = extract(html, final, article=p.article, anchor=l.anchor, selector=l.selector)
            res = Result(p, l, final_url=final, checked_at=stamp, availability=r["availability"] or "",
                         article_found=r["article_found"], method=r["method"] or "")
            if code in (403, 429) and r["price"] is None:
                r["blocked"] = r["blocked"] or f"HTTP {code}"
            if r["price"] is not None:
                res.status, res.price = "ok", r["price"]
            elif r["blocked"]:
                res.status, res.error = "blocked", f"Защита от ботов («{r['blocked']}») — не обходим"
                blocked_sites[site] = r["blocked"]
            elif r["not_found"] or code == 404:
                res.status, res.error = "notfound", f"HTTP {code}" if code == 404 else "Похоже, товар удалён"
            else:
                res.status = "noprice"
                res.error = (f"На странице нет текста «{l.anchor}»" if l.anchor and r["anchor_found"] is False
                             else ("Пустая страница — сайт строит её через JS (добавьте домен в use_browser_for)"
                                   if r["text_length"] < 200 else "Цена не найдена — задайте CSS-селектор"))
            results[id(l)] = res
            tag = (p.article or p.name)[:20]
            if res.status == "ok":
                d = deviation(res.price, p.our_price)
                dump = is_dumping(res.price, p.our_price, st.dump_threshold_pct)
                warn = " ⚠ артикул не найден на странице" if res.article_found is False else ""
                log(f"{stamp} [{pause:4.1f} с] {host_of(url)} · {tag}: {fmt_money(res.price)} ({fmt_pct(d)})"
                    f"{' — ДЕМПИНГ' if dump else ''}{warn}")
            else:
                log(f"{stamp} [{pause:4.1f} с] {host_of(url)} · {tag}: {STATUS_TEXT[res.status]} — {res.error}")

    try:
        for attempt, lst in ((0, queue), (1, retry)):
            for url, items in lst:
                if stop is not None and stop.is_set():
                    log("Остановлено пользователем.")
                    break
                process(url, items, attempt)
    finally:
        if own_fetcher:
            fetcher.close()
    for p in positions:  # исключённые и выключенные — тоже в отчёт, чтобы было видно
        for l in p.links:
            if id(l) not in results:
                results[id(l)] = Result(p, l, status="excluded" if is_excluded(l.url, st.excluded_domains) else "skipped",
                                        error="домен в списке исключений" if is_excluded(l.url, st.excluded_domains) else "ссылка отключена",
                                        final_url=l.url)
    ordered = [results[id(l)] for p in positions for l in p.links]
    ok = sum(r.status == "ok" for r in ordered)
    dumps = sum(r.status == "ok" and bool(is_dumping(r.price, r.position.our_price, st.dump_threshold_pct)) for r in ordered)
    log(f"Готово за {time.time() - t_start:.0f} с: цен получено {ok} из {len(ordered)}, демпинг — {dumps}.")
    return ordered


def run_auto(positions: list, settings: Settings = None, log: Callable[[str], None] = print, **kw) -> list:
    """Для ноутбука. Без браузера — обычный run(). С Playwright (use_browser_for) — run() в отдельном потоке:
    sync-API Playwright не работает внутри asyncio-цикла Jupyter/Colab. Кнопка «стоп» (KeyboardInterrupt)
    останавливает проверку после текущей страницы, частичные результаты возвращаются."""
    import threading
    st = (settings or Settings()).normalized()
    stop = threading.Event()
    if not st.use_browser_for:
        try:
            return run(positions, st, log=log, stop=stop, **kw)
        except KeyboardInterrupt:
            log("Прервано.")
            raise
    box = {}
    sleep = kw.pop("sleep", None) or (lambda sec: stop.wait(sec))

    def work():
        try:
            box["res"] = run(positions, st, log=log, sleep=sleep, stop=stop, **kw)
        except BaseException as e:  # noqa: BLE001 — передаём в основной поток
            box["err"] = e

    t = threading.Thread(target=work, daemon=True)
    t.start()
    try:
        while t.is_alive():
            t.join(0.5)
    except KeyboardInterrupt:
        stop.set()
        log("Останавливаю после текущей страницы…")
        t.join()
    if "err" in box:
        raise box["err"]
    return box.get("res", [])


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# --------------------------------------------------------------------------------------------
# Отчёт
# --------------------------------------------------------------------------------------------
REPORT_COLUMNS = ["Товар", "Артикул", "Наша цена", "Цена товара", "Разница в цене в %", "Магазин демпинг",
                  "Ссылка на магазин конкурента", "Наличие", "Статус", "Проверено"]


def report_rows(results: list, settings: Settings = None) -> list:
    st = (settings or Settings()).normalized()
    rows = []
    for r in results:
        if r.status == "excluded":
            continue
        p = r.position
        d = deviation(r.price, p.our_price) if r.status == "ok" else None
        dump = is_dumping(r.price, p.our_price, st.dump_threshold_pct) if r.status == "ok" else None
        store = host_of(r.link.url)
        status = STATUS_TEXT.get(r.status, r.status) + (f": {r.error}" if r.error and r.status != "ok" else "")
        if r.status == "ok" and r.article_found is False:
            status += " · ⚠ артикул не найден на странице"
        rows.append({
            "Товар": p.name, "Артикул": p.article, "Наша цена": p.our_price or None,
            "Цена товара": r.price if r.status == "ok" else None,
            "Разница в цене в %": d,
            "Магазин демпинг": f"{store} — " + ("Да" if dump else "Нет" if dump is not None else "н/д"),
            "Ссылка на магазин конкурента": r.final_url or r.link.url,
            "Наличие": r.availability, "Статус": status, "Проверено": r.checked_at,
        })
    rows.sort(key=lambda x: (x["Артикул"], x["Цена товара"] is None, x["Цена товара"] or 0))
    return rows


def report_df(results: list, settings: Settings = None):
    import pandas as pd
    return pd.DataFrame(report_rows(results, settings), columns=REPORT_COLUMNS)


def report_csv(results: list, settings: Settings = None) -> str:
    """CSV для русского Excel: BOM, «;», десятичная запятая."""
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(REPORT_COLUMNS)
    for row in report_rows(results, settings):
        w.writerow(["" if row[c] is None else (str(row[c]).replace(".", ",") if c == "Разница в цене в %" else row[c])
                    for c in REPORT_COLUMNS])
    return "\ufeff" + buf.getvalue()


def save_xlsx(results: list, path: str, settings: Settings = None) -> str:
    """Excel с подсветкой демпинга и листом «Сводка» (минимальная цена по каждой позиции)."""
    import pandas as pd
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    df = report_df(results, settings)
    st = (settings or Settings()).normalized()
    summary = []
    for p in {id(r.position): r.position for r in results}.values():
        prices = [(r.price, host_of(r.link.url)) for r in results if r.position is p and r.status == "ok"]
        mn = min(prices) if prices else (None, "")
        summary.append({"Товар": p.name, "Артикул": p.article, "Наша цена": p.our_price,
                        "Мин. цена конкурента": mn[0], "Где": mn[1], "Разница в %": deviation(mn[0], p.our_price),
                        "Магазинов с демпингом": sum(bool(is_dumping(x[0], p.our_price, st.dump_threshold_pct)) for x in prices),
                        "Цен получено": len(prices)})
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="Конкуренты", index=False)
        pd.DataFrame(summary).to_excel(xw, sheet_name="Сводка", index=False)
        red = PatternFill("solid", fgColor="FCE8E6")
        for name, frame in (("Конкуренты", df), ("Сводка", pd.DataFrame(summary))):
            ws = xw.sheets[name]
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for i, col in enumerate(frame.columns, 1):
                width = max([len(str(col))] + [len(str(v)) for v in frame[col].head(200)]) + 2
                ws.column_dimensions[get_column_letter(i)].width = min(60, width)
            ws.freeze_panes = "A2"
        ws = xw.sheets["Конкуренты"]
        ci = REPORT_COLUMNS.index("Магазин демпинг") + 1
        for row in ws.iter_rows(min_row=2):
            if str(row[ci - 1].value or "").endswith("— Да"):
                for c in row:
                    c.fill = red
    return path


def styled(results: list, settings: Settings = None):
    """Таблица для Jupyter/Colab: красным — демпинг."""
    df = report_df(results, settings)

    def hl(row):
        return ["background-color:#fce8e6" if str(row["Магазин демпинг"]).endswith("— Да") else ""] * len(row)

    try:
        import jinja2  # noqa: F401 — нужен для df.style (в Colab есть)
    except ImportError:
        return df
    return (df.style.apply(hl, axis=1)
            .format({"Наша цена": lambda v: fmt_money(v) if v == v and v is not None else "—",
                     "Цена товара": lambda v: fmt_money(v) if v == v and v is not None else "—",
                     "Разница в цене в %": lambda v: fmt_pct(v) if v == v and v is not None else "—"})
            .hide(axis="index"))


def history_rows(results: list, settings: Settings = None) -> list:
    """Строки для накопительной истории (одна строка = одна цена в момент проверки)."""
    return [{"Проверено": r["Проверено"], "Артикул": r["Артикул"], "Магазин": r["Магазин демпинг"].split(" — ")[0],
             "Цена": r["Цена товара"], "Наша цена": r["Наша цена"], "Разница в %": r["Разница в цене в %"],
             "Ссылка": r["Ссылка на магазин конкурента"]}
            for r in report_rows(results, settings) if r["Цена товара"] is not None]
